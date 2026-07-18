import shutil
from pathlib import Path

import openpyxl
from docx import Document
from pypdf import PdfReader

from kf.config import Settings
from kf.ocr import extract_text_from_image
from kf.transcribe import transcribe_audio
from kf.video import extract_audio, sample_frames
from kf.vision_caption import caption_image

PLAIN_TEXT_EXTENSIONS = {".md", ".txt", ".csv", ".html"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".webm"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".ogg", ".m4a"}


def _extract_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text)


def _extract_xlsx(path: Path) -> str:
    workbook = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"[Лист: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_image(path: Path, settings: Settings) -> str:
    ocr_text = extract_text_from_image(path, languages=settings.ocr_languages)
    if len(ocr_text.strip()) < settings.image_caption_threshold_chars:
        caption = caption_image(settings, path)
        return f"{ocr_text}\n\n[Описание изображения]\n{caption}".strip()
    return ocr_text


def _extract_audio_file(path: Path, settings: Settings) -> str:
    return transcribe_audio(path, settings.whisper_model_size, settings.model_cache_dir)


def _extract_video(path: Path, settings: Settings) -> str:
    try:
        audio_path = extract_audio(path)
    except Exception:
        audio_path = None

    frames_dir = None
    try:
        if audio_path is not None:
            transcript = transcribe_audio(
                audio_path, model_size=settings.whisper_model_size, cache_dir=settings.model_cache_dir
            )
        else:
            transcript = ""

        frames_dir, frames = sample_frames(
            path,
            interval_seconds=settings.video_frame_interval_seconds,
            max_frames=settings.max_video_frames,
        )
        frame_blocks = []
        for i, frame_path in enumerate(frames):
            timestamp_seconds = i * settings.video_frame_interval_seconds
            minutes, seconds = divmod(timestamp_seconds, 60)
            frame_text = _extract_image(frame_path, settings)
            frame_blocks.append(f"[Кадр {minutes:02d}:{seconds:02d}]\n{frame_text}")

        parts = ([f"[Транскрипт]\n{transcript}"] if transcript else []) + frame_blocks
        return "\n\n".join(parts).strip()
    finally:
        if audio_path is not None:
            Path(audio_path).unlink(missing_ok=True)
        if frames_dir is not None:
            shutil.rmtree(frames_dir, ignore_errors=True)


def extract_text(path: Path, settings: Settings) -> str:
    suffix = path.suffix.lower()
    if suffix in PLAIN_TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in IMAGE_EXTENSIONS:
        return _extract_image(path, settings)
    if suffix in VIDEO_EXTENSIONS:
        return _extract_video(path, settings)
    if suffix in AUDIO_EXTENSIONS:
        return _extract_audio_file(path, settings)
    if suffix == ".xlsx":
        return _extract_xlsx(path)
    raise ValueError(f"Unsupported file type: {suffix}")
